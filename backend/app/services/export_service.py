"""结果导出：统计报表、预算、借款、报销、获奖台账导出为 Excel / Word / CSV。"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import settings
from ..db.models import Award, Budget, Loan, Registration, Reimbursement
from .stats_service import StatsFilter, StatsService


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d-%H%M%S")


def _safe(name: str) -> str:
    return "".join(c for c in name if c not in r'\/:*?"<>|').strip() or "export"


def export_path(filename: str) -> Path:
    return settings.export_path / _safe(filename)


# ============================================================
#  通用写出
# ============================================================
def to_excel_bytes(sheets: dict[str, pd.DataFrame], *, title: str = "") -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        if not sheets:
            pd.DataFrame({"提示": ["暂无数据"]}).to_excel(writer, sheet_name="说明", index=False)
        for name, df in sheets.items():
            sheet_name = _safe(name)[:31] or "Sheet"
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            _autofit(writer.sheets[sheet_name], df)
    return buf.getvalue()


def _autofit(worksheet, df: pd.DataFrame) -> None:
    try:
        for idx, column in enumerate(df.columns, start=1):
            max_len = max([len(str(column))] + [len(str(v)) for v in df[column].head(200).tolist()] or [8])
            width = min(max(max_len + 2, 8), 42)
            worksheet.column_dimensions[worksheet.cell(row=1, column=idx).column_letter].width = width
        worksheet.freeze_panes = "A2"
    except Exception:  # noqa: BLE001 - 样式失败不影响导出内容
        pass


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    text = df.to_csv(index=False, lineterminator="\n")
    return ("\ufeff" + text).encode("utf-8")  # BOM 便于 Excel 直接打开


def to_docx_bytes(title: str, body_text: str, tables: list[tuple[str, pd.DataFrame]] | None = None) -> bytes:
    doc = DocxDocument()

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    for line in (body_text or "").split("\n"):
        para = doc.add_paragraph(line)
        para.paragraph_format.space_after = Pt(2)

    for caption, df in tables or []:
        doc.add_paragraph("")
        cap = doc.add_paragraph()
        cap.add_run(caption).bold = True
        table = doc.add_table(rows=1, cols=len(df.columns))
        table.style = "Table Grid"
        for i, col in enumerate(df.columns):
            cell = table.rows[0].cells[i]
            cell.text = str(col)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
        for _, row in df.iterrows():
            cells = table.add_row().cells
            for i, col in enumerate(df.columns):
                cells[i].text = str(row[col])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ============================================================
#  统计报表
# ============================================================
class ExportService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.stats = StatsService(db)

    # ---------- 统计报表 ----------
    def build_stats_sheets(self, flt: StatsFilter | None = None) -> dict[str, pd.DataFrame]:
        flt = flt or StatsFilter()
        overview = self.stats.overview(flt)
        sheets: dict[str, pd.DataFrame] = {
            "统计口径": pd.DataFrame(
                [{"项目": k, "内容": str(v)} for k, v in overview.items() if k in ("filter", "data_source")]
            ),
            "总览": pd.DataFrame(
                [
                    {"指标": "参赛队伍数", "数值": overview["registrations"]},
                    {"指标": "参赛人次", "数值": overview["participants"]},
                    {"指标": "获奖数量", "数值": overview["awards"]},
                    {"指标": "获奖率(%)", "数值": overview["award_rate"]},
                    {"指标": "高等级奖项数（特等+一等）", "数值": overview["high_level_awards"]},
                    {"指标": "参与学生数", "数值": overview["students"]},
                    {"指标": "指导教师数", "数值": overview["teachers"]},
                    {"指标": "涉及赛事数", "数值": overview["competitions"]},
                ]
            ),
            "年度统计": pd.DataFrame(self.stats.by_year(flt)),
            "赛事统计": pd.DataFrame(self.stats.by_competition(flt)),
            "专业统计": pd.DataFrame(self.stats.by_major(flt)),
            "获奖等级统计": pd.DataFrame(self.stats.by_award_level(flt)),
            "赛事级别统计": pd.DataFrame(self.stats.by_level(flt)),
            "教师指导统计": pd.DataFrame(self.stats.by_teacher(flt)),
            "教师工作量": pd.DataFrame(
                [
                    {k: v for k, v in row.items() if k != "detail"}
                    for row in self.stats.teacher_workload(flt)
                ]
            ),
            "学生获奖统计": pd.DataFrame(self.stats.by_student(flt)),
        }
        return {k: v for k, v in sheets.items() if v is not None and not v.empty}

    def stats_excel(self, flt: StatsFilter | None = None) -> tuple[str, bytes]:
        sheets = self.build_stats_sheets(flt)
        name = f"竞赛统计报表-{_stamp()}.xlsx"
        return name, to_excel_bytes(sheets)

    # ---------- 获奖台账 ----------
    def awards_sheets(self, flt: StatsFilter | None = None) -> dict[str, pd.DataFrame]:
        flt = flt or StatsFilter()
        _, df_aw = self.stats.load_frames(flt)
        if df_aw.empty:
            return {"获奖明细": pd.DataFrame({"提示": ["暂无符合条件的数据"]})}
        detail = df_aw.rename(
            columns={
                "year": "年度",
                "competition_name": "赛事名称",
                "track_name": "赛道",
                "award_level": "获奖等级",
                "award_grade": "级别",
                "project_title": "作品/项目名称",
                "team_name": "队伍名称",
                "students": "获奖学生",
                "majors": "专业",
                "mentors": "指导教师",
                "points": "折算积分",
                "award_date": "获奖日期",
            }
        )[
            ["年度", "赛事名称", "赛道", "级别", "获奖等级", "作品/项目名称", "队伍名称", "获奖学生", "专业", "指导教师", "折算积分", "获奖日期"]
        ]
        return {
            "获奖明细": detail,
            "按年度汇总": pd.DataFrame(self.stats.by_year(flt)),
            "按赛事汇总": pd.DataFrame(self.stats.by_competition(flt)),
            "按专业汇总": pd.DataFrame(self.stats.by_major(flt)),
        }

    def awards_excel(self, flt: StatsFilter | None = None) -> tuple[str, bytes]:
        return f"竞赛获奖台账-{_stamp()}.xlsx", to_excel_bytes(self.awards_sheets(flt))

    # ---------- 参赛台账 ----------
    def registrations_excel(self, flt: StatsFilter | None = None) -> tuple[str, bytes]:
        flt = flt or StatsFilter()
        df_reg, _ = self.stats.load_frames(flt)
        if df_reg.empty:
            df = pd.DataFrame({"提示": ["暂无符合条件的数据"]})
        else:
            df = df_reg.rename(
                columns={
                    "year": "年度",
                    "competition_name": "赛事名称",
                    "track_name": "赛道",
                    "competition_level": "级别",
                    "competition_category": "类别",
                    "team_name": "队伍名称",
                    "project_title": "作品/项目名称",
                    "status": "状态",
                    "stage_reached": "晋级阶段",
                    "lead_major": "队长专业",
                    "mentor_names": "指导教师",
                    "member_count": "队员人数",
                    "members": "队员",
                }
            )[
                ["年度", "赛事名称", "赛道", "级别", "类别", "队伍名称", "作品/项目名称", "状态", "晋级阶段",
                 "队长专业", "指导教师", "队员人数", "队员"]
            ]
        return f"竞赛参赛台账-{_stamp()}.xlsx", to_excel_bytes({"参赛明细": df})

    # ---------- 预算 / 借款 / 报销 ----------
    def budget_doc(self, budget: Budget, *, as_word: bool = True) -> tuple[str, bytes, str]:
        items = [
            {
                "序号": i,
                "费用类别": it.category,
                "项目": it.item_name,
                "单位": it.unit,
                "数量": it.quantity,
                "单价(元)": it.unit_price,
                "金额(元)": float(it.amount),
                "费用标准编号": it.standard_code,
                "计算依据": it.calc_basis,
                "备注": it.remark,
            }
            for i, it in enumerate(budget.items, start=1)
        ]
        df = pd.DataFrame(items)
        if as_word:
            tables = [("预算明细", df)] if items else []
            data = to_docx_bytes(budget.title or "竞赛经费预算申请", budget.generated_text or "", tables)
            return f"{budget.doc_no}-预算申请.docx", data, "docx"
        sheets = {
            "预算明细": df if items else pd.DataFrame({"提示": ["无明细"]}),
            "预算申请正文": pd.DataFrame({"正文": (budget.generated_text or "").split("\n")}),
        }
        return f"{budget.doc_no}-预算申请.xlsx", to_excel_bytes(sheets), "xlsx"

    def loan_doc(self, loan: Loan, *, as_word: bool = True) -> tuple[str, bytes, str]:
        content = loan.generated_content or {}
        df = pd.DataFrame([{"字段": k, "内容": str(v)} for k, v in content.items()])
        checks = pd.DataFrame(loan.checks or [])
        if as_word:
            labels = {"pass": "通过", "error": "必须修改", "warning": "建议修改", "info": "提示"}
            lines = ["借 款 单", ""]
            for key, value in content.items():
                if value not in ("", None):
                    lines.append(f"{key}：{value}")
            if loan.checks:
                lines += ["", "附：事前校验结果"]
                for c in loan.checks:
                    lines.append(f"· [{labels.get(c.get('level'), c.get('level', ''))}] {c.get('message', '')}")
                    if c.get("suggestion"):
                        lines.append(f"  建议：{c['suggestion']}")
            data = to_docx_bytes(
                "借款单",
                "\n".join(lines),
                [("借款单要素", df)] if not df.empty else [],
            )
            return f"{loan.doc_no}-借款单.docx", data, "docx"
        sheets = {"借款单要素": df, "校验结果": checks if not checks.empty else pd.DataFrame({"提示": ["无"]})}
        return f"{loan.doc_no}-借款单.xlsx", to_excel_bytes(sheets), "xlsx"

    def reimbursement_doc(self, reimb: Reimbursement, *, as_word: bool = True) -> tuple[str, bytes, str]:
        items = [
            {
                "序号": i,
                "费用类别": it.category,
                "项目": it.item_name,
                "金额(元)": float(it.amount),
                "发生日期": str(it.expense_date) if it.expense_date else "",
                "发票号": it.invoice_no,
                "是否有票据": "是" if it.has_invoice else "否",
                "计算依据": it.calc_basis,
                "备注": it.remark,
            }
            for i, it in enumerate(reimb.items, start=1)
        ]
        df = pd.DataFrame(items)
        checks = pd.DataFrame(reimb.checks or [])
        if as_word:
            tables = []
            if not df.empty:
                tables.append(("报销明细", df))
            if not checks.empty:
                checks_display = checks.rename(columns={"code": "检查项", "level": "结论", "message": "说明", "suggestion": "建议"})
                tables.append(("材料完整性校验", checks_display))
            data = to_docx_bytes(
                "竞赛经费报销单", reimb.generated_text or "", tables
            )
            return f"{reimb.doc_no}-报销单.docx", data, "docx"
        sheets = {
            "报销明细": df if not df.empty else pd.DataFrame({"提示": ["无明细"]}),
            "材料完整性校验": checks if not checks.empty else pd.DataFrame({"提示": ["无"]}),
        }
        return f"{reimb.doc_no}-报销单.xlsx", to_excel_bytes(sheets), "xlsx"

    # ---------- 材料检查报告 ----------
    def review_doc(self, title: str, payload: dict[str, Any]) -> tuple[str, bytes, str]:
        lines = [title, ""]
        for section in payload.get("sections", []):
            lines.append(f"【{section.get('name')}】")
            lines.append(str(section.get("summary") or ""))
            for point in section.get("points", []):
                lines.append(f"· {point}")
            lines.append("")
        if payload.get("questions"):
            lines.append("【预答辩问题】")
            for i, q in enumerate(payload["questions"], start=1):
                lines.append(f"{i}. {q}")
        tables = []
        if payload.get("rubric"):
            tables.append(("评审量表", pd.DataFrame(payload["rubric"])))
        data = to_docx_bytes(title, "\n".join(lines), tables)
        return f"材料检查报告-{_stamp()}.docx", data, "docx"


# ============================================================
#  明细查询（供前端导出中心）
# ============================================================
def list_budgets(db: Session, limit: int = 100) -> list[dict[str, Any]]:
    rows = db.execute(select(Budget).order_by(Budget.created_at.desc()).limit(limit)).scalars()
    return [_budget_row(b) for b in rows]


def list_loans(db: Session, limit: int = 100) -> list[dict[str, Any]]:
    rows = db.execute(select(Loan).order_by(Loan.created_at.desc()).limit(limit)).scalars()
    return [
        {
            "id": l.id,
            "doc_no": l.doc_no,
            "borrower_name": l.borrower_name,
            "amount": float(l.amount or 0),
            "amount_upper": l.amount_upper,
            "purpose": l.purpose,
            "competition_name": l.competition_name,
            "status": l.status,
            "expected_repay_date": str(l.expected_repay_date) if l.expected_repay_date else "",
            "created_at": l.created_at.isoformat() if l.created_at else "",
        }
        for l in rows
    ]


def list_reimbursements(db: Session, limit: int = 100) -> list[dict[str, Any]]:
    rows = db.execute(select(Reimbursement).order_by(Reimbursement.created_at.desc()).limit(limit)).scalars()
    return [
        {
            "id": r.id,
            "doc_no": r.doc_no,
            "claimant_name": r.claimant_name,
            "competition_name": r.competition_name,
            "total_amount": float(r.total_amount or 0),
            "approved_budget_amount": float(r.approved_budget_amount) if r.approved_budget_amount is not None else None,
            "loan_amount": float(r.loan_amount) if r.loan_amount is not None else None,
            "status": r.status,
            "completeness_score": r.completeness_score,
            "created_at": r.created_at.isoformat() if r.created_at else "",
        }
        for r in rows
    ]


def _budget_row(b: Budget) -> dict[str, Any]:
    return {
        "id": b.id,
        "doc_no": b.doc_no,
        "title": b.title,
        "competition_name": b.competition_name,
        "year": b.year,
        "applicant_name": b.applicant_name,
        "team_name": b.team_name,
        "student_count": b.student_count,
        "teacher_count": b.teacher_count,
        "destination": b.destination,
        "trip_days": b.trip_days,
        "total_amount": float(b.total_amount or 0),
        "approved_amount": float(b.approved_amount) if b.approved_amount is not None else None,
        "status": b.status,
        "created_at": b.created_at.isoformat() if b.created_at else "",
        "item_count": len(b.items),
    }